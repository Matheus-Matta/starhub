"""Relatorios do menu: um por model do tema, periodo da criacao, Excel e PDF."""

from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO

import pytest
from django.contrib.auth.models import Permission
from django.core.management import call_command
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.models import User
from apps.core.relatorios import exportar, registro
from apps.loja.models import Cliente, Pedido, Tag

pytestmark = pytest.mark.django_db
PERIODO = "?de=2026-09-01&ate=2026-09-30"


def _url(chave, query=PERIODO):
    return reverse("admin:relatorio", args=[chave]) + query


def _menu(cliente):
    html = cliente.get(reverse("admin:index")).content.decode()
    return html.split('class="sidebar-menu"')[1].split("</nav>")[0]


@pytest.fixture
def tags():
    dentro, fora = Tag.objects.create(nome="Setembro"), Tag.objects.create(nome="Janeiro")
    Tag.objects.filter(pk=dentro.pk).update(
        created_at=timezone.make_aware(datetime(2026, 9, 30, 23, 30)))
    Tag.objects.filter(pk=fora.pk).update(
        created_at=timezone.make_aware(datetime(2026, 1, 10, 12)))
    return dentro, fora


def test_menu_tem_dropdown_de_relatorios_com_um_item_por_model(admin_logado):
    menu = _menu(admin_logado)
    grupo = menu.split('data-menu-grupo="relatorios"')[1].split("</details>")[0]

    for chave in ("uso-cupons", "loja.pedido", "loja.cupom", "loja.tag", "core.user"):
        assert reverse("admin:relatorio", args=[chave]) in grupo
    # Configuracao unica por conta nao tem lista para relatar.
    assert "integracoes.configuracaointegracao" not in grupo
    assert "notificacoes.configuracaoemail" not in grupo


def test_menu_so_mostra_relatorio_do_que_a_pessoa_pode_ver(client, conta):
    pessoa = User.objects.create_user("vendas", "v@a.test", "senha-forte-123",
                                      is_staff=True, account=conta)
    pessoa.user_permissions.add(Permission.objects.get(codename="view_tag"))
    client.force_login(pessoa)

    grupo = _menu(client).split('data-menu-grupo="relatorios"')[1].split("</details>")[0]

    assert reverse("admin:relatorio", args=["loja.tag"]) in grupo
    assert reverse("admin:relatorio", args=["loja.pedido"]) not in grupo
    assert reverse("admin:relatorio", args=["uso-cupons"]) not in grupo
    assert client.get(_url("loja.pedido")).status_code == 403


def test_pagina_mostra_so_o_periodo_e_inclui_o_ultimo_dia_inteiro(admin_logado, tags):
    resposta = admin_logado.get(_url("loja.tag"))

    assert resposta.status_code == 200
    html = resposta.content.decode()
    assert "Setembro" in html and "Janeiro" not in html
    assert "1 registro" in html


def test_excel_tem_as_colunas_da_lista_e_as_linhas_do_periodo(admin_logado, tags):
    resposta = admin_logado.get(_url("loja.tag", PERIODO + "&formato=xlsx"))

    assert resposta.status_code == 200
    assert resposta["Content-Type"] == exportar.TIPOS["xlsx"]
    assert "attachment" in resposta["Content-Disposition"]
    linhas = list(load_workbook(BytesIO(resposta.content)).active.values)
    assert "Nome" in linhas[0]
    assert len(linhas) == 2 and "Setembro" in linhas[1]


def test_pdf_sai_como_arquivo_pdf(admin_logado, tags):
    resposta = admin_logado.get(_url("loja.tag", PERIODO + "&formato=pdf"))

    assert resposta.status_code == 200
    assert resposta["Content-Type"] == "application/pdf"
    assert resposta.content.startswith(b"%PDF")


def test_acima_do_limite_nao_gera_arquivo_e_pede_periodo_menor(admin_logado, tags, monkeypatch):
    """Exportacao sem teto prende o worker do site gerando um PDF de milhares de paginas."""
    monkeypatch.setitem(exportar.LIMITES, "pdf", 0)

    resposta = admin_logado.get(_url("loja.tag", PERIODO + "&formato=pdf"), follow=True)

    assert resposta.redirect_chain
    assert "Escolha um periodo menor" in resposta.content.decode()


def test_todo_relatorio_abre_e_exporta_com_dados_reais(admin_logado, conta, settings, tmp_path):
    """Coluna da lista que so funciona em HTML (badge, imagem, link) nao pode
    derrubar o Excel nem o PDF de nenhum model."""
    settings.MEDIA_ROOT = tmp_path
    call_command("gerar_demo", "--conta", conta.slug, "--permitir-producao", stdout=StringIO())

    com_dados = 0
    for relatorio in registro.todos():
        respostas = [admin_logado.get(_url(relatorio.chave, "?de=&ate=" + formato))
                     for formato in ("", "&formato=xlsx", "&formato=pdf")]
        assert [r.status_code for r in respostas] == [200] * 3, relatorio.chave
        total = respostas[0].context["total"]
        planilha = load_workbook(BytesIO(respostas[1].content)).active
        assert planilha.max_row - 1 == total, relatorio.chave
        com_dados += total > 0
    # O demo cobre a maior parte dos models: sem isso o laco acima testaria tabela vazia.
    assert com_dados >= 10


def test_coluna_com_avatar_exporta_so_o_dado_sem_as_iniciais(admin_logado):
    """A coluna Cliente da lista tem avatar com iniciais: no Excel saia "AL Ana Lima"."""
    cliente = Cliente.objects.create(email="ana@x.com", nome="Ana", sobrenome="Lima")
    Pedido.objects.create(number="1001", cliente=cliente)

    resposta = admin_logado.get(_url("loja.pedido", "?de=&ate=&formato=xlsx"))

    cabecalho, linha = list(load_workbook(BytesIO(resposta.content)).active.values)
    assert linha[cabecalho.index("Cliente")] == "Ana Lima ana@x.com"


def test_coluna_de_valor_formatada_na_tela_vira_numero_no_excel(admin_logado):
    """"R$ 1.347,70" como texto nao soma na planilha; o campo por tras e o total."""
    Pedido.objects.create(number="1001", total=Decimal("1347.70"))

    planilha = load_workbook(BytesIO(admin_logado.get(
        _url("loja.pedido", "?de=&ate=&formato=xlsx")).content)).active
    cabecalho = [celula.value for celula in planilha[1]]
    total = planilha.cell(row=2, column=cabecalho.index("Total") + 1)

    assert Decimal(str(total.value)) == Decimal("1347.70")
    assert "R$" in total.number_format


def test_relatorio_desconhecido_e_404(admin_logado):
    assert admin_logado.get(_url("loja.naoexiste")).status_code == 404
