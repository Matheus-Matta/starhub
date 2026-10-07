"""Arquivos dos relatorios: Excel (openpyxl) e PDF (reportlab), mais o texto da tela.

O Excel recebe numero e data de verdade (soma e filtro funcionam na planilha); o PDF
e a tela usam o texto no formato brasileiro.
"""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font

from apps.core.relatorios import pdf
from apps.core.ui.formatos import moeda

TIPOS = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
# Teto de linhas por arquivo: o arquivo e gerado dentro da requisicao, e um PDF de
# milhares de paginas prenderia o processo do site. Acima disso, periodo menor.
LIMITES = {"xlsx": 50000, "pdf": 5000}
_FORMATO_EXCEL = {datetime: "dd/mm/yyyy hh:mm", date: "dd/mm/yyyy"}
_PROIBIDO_NA_ABA = str.maketrans("", "", "[]:*?/\\")


def texto(valor, em_reais=False):
    """Celula como o brasileiro le: 1.234,50 / 30/09/2026 14:05."""
    if valor is None:
        return ""
    if isinstance(valor, Decimal):
        if em_reais:
            return moeda(valor)
        # Sem quantizar: peso e medida guardam as casas que tem.
        return f"{valor:,f}".replace(",", "_").replace(".", ",").replace("_", ".")
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return str(valor)


def _aba(titulo):
    return str(titulo).translate(_PROIBIDO_NA_ABA)[:31] or "Relatorio"


def _preencher(aba, tabela):
    aba.append(tabela.colunas)
    for celula in aba[1]:
        celula.font = Font(bold=True)
    larguras = [len(str(coluna)) for coluna in tabela.colunas]
    for linha in tabela.linhas:
        aba.append(linha)
        for indice, valor in enumerate(linha):
            larguras[indice] = max(larguras[indice], len(texto(valor, indice in tabela.moeda)))
    for indice, celula in enumerate(aba[1], start=0):
        aba.column_dimensions[celula.column_letter].width = min(larguras[indice] + 2, 50)
    for linha in aba.iter_rows(min_row=2):
        for celula in linha:
            _formatar(celula, celula.col_idx - 1 in tabela.moeda)
    aba.freeze_panes = "A2"


def _formatar(celula, em_reais):
    valor = celula.value
    if isinstance(valor, Decimal):
        celula.number_format = '"R$" #,##0.00' if em_reais else "#,##0.00"
    elif isinstance(valor, datetime):
        celula.number_format = _FORMATO_EXCEL[datetime]
    elif isinstance(valor, date):
        celula.number_format = _FORMATO_EXCEL[date]


def xlsx(resultado):
    livro = Workbook()
    _preencher(livro.active, resultado.principal)
    livro.active.title = _aba(resultado.principal.titulo)
    for resumo in resultado.resumos:
        _preencher(livro.create_sheet(_aba(resumo.titulo)), resumo)
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def descrever_periodo(de, ate):
    if de and ate:
        return f"Periodo: {texto(de)} a {texto(ate)}"
    if de or ate:
        return f"Periodo: {'desde ' + texto(de) if de else 'ate ' + texto(ate)}"
    return "Periodo: todos os registros"


def resposta(relatorio, resultado, formato, de, ate):
    periodo = descrever_periodo(de, ate)
    if formato == "pdf":
        conteudo = pdf.gerar(relatorio.titulo, periodo, resultado, texto)
    else:
        conteudo = xlsx(resultado)
    sufixo = "-".join(d.isoformat() for d in (de, ate) if d) or "tudo"
    saida = HttpResponse(conteudo, content_type=TIPOS[formato])
    saida["Content-Disposition"] = (
        f'attachment; filename="{relatorio.nome_arquivo}-{sufixo}.{formato}"')
    return saida
