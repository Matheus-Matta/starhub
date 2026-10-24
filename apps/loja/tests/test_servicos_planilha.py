"""Planilha de servicos: exportar, editar e importar de volta sem duplicar."""

from decimal import Decimal
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.integracoes.importar_planilha import LinhaInvalida
from apps.integracoes.models import ExecucaoIntegracao
from apps.loja.models import Categoria, RegraServico, Servico
from apps.loja.services.servicos_planilha import COLUNAS, importar_linha, linhas
from apps.loja.services.variantes import criar_produto

pytestmark = pytest.mark.django_db
NOMES = [coluna for coluna, _, _ in COLUNAS]


@pytest.fixture
def catalogo():
    sofas = Categoria.objects.create(nome="Sofas", slug="sofas")
    sofa = criar_produto("Sofa", sku="SOF-1", price=Decimal("1500.00"))
    mesa = criar_produto("Mesa", sku="MES-1", price=Decimal("800.00"))
    return sofas, sofa, mesa


def _linha(**campos):
    return {"sku": "SERV-MONT", "nome": "Montagem", "preco": "120,00", **campos}


def test_linha_cria_servico_e_regra_por_categoria_com_faixa(catalogo):
    sofas, _, _ = catalogo
    servico, criado = importar_linha(_linha(regra_ordem="10", regra_tipo="categoria",
                                            regra_preco="80,00", categorias="Sofas",
                                            preco_ate="2.000,00"))

    regra = servico.regras.get()
    assert criado and servico.preco == Decimal("120.00")
    assert (regra.tipo, regra.preco, regra.preco_ate) == (
        "categoria", Decimal("80.00"), Decimal("2000.00"))
    assert list(regra.categorias.all()) == [sofas]


def test_importar_de_novo_atualiza_pelo_sku_e_pela_ordem_sem_duplicar(catalogo):
    _, sofa, mesa = catalogo
    importar_linha(_linha(regra_ordem="10", regra_tipo="produto", produtos="SOF-1"))

    servico, criado = importar_linha(_linha(preco="150,00", regra_ordem="10",
                                            regra_tipo="produto", produtos="SOF-1|MES-1"))

    assert not criado and Servico.objects.count() == 1 and RegraServico.objects.count() == 1
    assert servico.preco == Decimal("150.00")
    assert set(servico.regras.get().produtos.all()) == {sofa, mesa}


@pytest.mark.parametrize(("campos", "motivo"), [
    ({"sku": "", "nome": ""}, "sku e nome vazios"),
    ({"sku": "NOVO", "nome": ""}, "preencha o nome"),
    ({"regra_ordem": "10", "regra_tipo": "categoria", "categorias": "Nao existe"},
     "categoria 'Nao existe' nao existe"),
    ({"regra_ordem": "10", "regra_tipo": "produto", "produtos": "XYZ"},
     "produto 'XYZ' nao existe"),
    ({"regra_ordem": "10", "regra_tipo": "produto"}, "ao menos um produto"),
    ({"regra_ordem": "10", "regra_tipo": "faixa"}, "categoria ou produto"),
    ({"preco": "abc"}, "preco 'abc' invalido"),
    ({"ativo": "talvez"}, "use sim ou nao"),
])
def test_linha_ruim_diz_o_que_corrigir_e_nao_grava_nada(catalogo, campos, motivo):
    with pytest.raises(LinhaInvalida, match=motivo):
        importar_linha(_linha(**campos))

    assert not RegraServico.objects.exists()


def test_exportado_volta_igual_na_importacao(catalogo):
    """Exportar, editar e importar e o caminho: a exportacao precisa ser lida de volta."""
    importar_linha(_linha(regra_ordem="10", regra_tipo="categoria", categorias="Sofas"))
    importar_linha(_linha(regra_ordem="20", regra_tipo="produto", produtos="MES-1",
                          regra_ativa="nao"))
    importar_linha({"sku": "SERV-GAR", "nome": "Garantia", "preco": "99,90"})
    exportadas = [dict(zip(NOMES, linha, strict=True)) for linha in linhas()]
    RegraServico.objects.all().delete()

    for dados in exportadas:
        importar_linha({k: "" if v is None else str(v) for k, v in dados.items()})

    assert [(e["sku"], e["regra_ordem"]) for e in exportadas] == [
        ("SERV-GAR", ""), ("SERV-MONT", 10), ("SERV-MONT", 20)]
    assert Servico.objects.count() == 2 and RegraServico.objects.count() == 2
    assert not RegraServico.objects.get(ordem=20).active


def test_tela_tem_importar_e_exportar_planilha_e_nao_importa_da_shopify(admin_logado):
    importar_linha(_linha())
    html = admin_logado.get(reverse("admin:loja_servico_changelist")).content.decode()

    assert reverse("admin:loja_servico_importar") in html
    assert reverse("admin:loja_servico_exportar") in html
    assert "Importar da Shopify" not in html

    resposta = admin_logado.get(reverse("admin:loja_servico_exportar"))
    aba = load_workbook(BytesIO(resposta.content)).active
    assert [c.value for c in aba[1]] == NOMES and aba["A2"].value == "SERV-MONT"


def test_aba_produtos_mostra_cada_produto_com_sku_preco_e_o_preco_do_servico(catalogo):
    """Pedido do operador: ver na planilha quais produtos cada regra pega e por quanto.
    O sofa cai nas duas regras; vale a de menor ordem, como no checkout."""
    sofas, sofa, mesa = catalogo
    sofa.categorias.add(sofas)
    importar_linha(_linha(regra_ordem="10", regra_tipo="produto", produtos="SOF-1",
                          regra_preco="80,00"))
    importar_linha(_linha(regra_ordem="20", regra_tipo="categoria", categorias="Sofas"))
    importar_linha(_linha(regra_ordem="30", regra_tipo="produto", produtos="MES-1",
                          regra_ativa="nao"))
    livro = load_workbook(BytesIO(_xlsx()))

    produtos = [[c.value for c in linha] for linha in livro["Produtos"].iter_rows(min_row=2)]

    assert livro.sheetnames == ["Servicos", "Produtos"]
    assert produtos == [["Montagem", "SERV-MONT", "sim", 10, "produto", "Sofa", "SOF-1",
                         1500, "Sofas", 80]]


def _xlsx():
    from apps.loja.services.servicos_planilha import planilha

    arquivo = BytesIO()
    planilha().save(arquivo)
    return arquivo.getvalue()


def test_importar_planilha_vira_tarefa(admin_logado, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    from django.core.files.uploadedfile import SimpleUploadedFile

    arquivo = SimpleUploadedFile("servicos.csv", "sku;nome\nSERV-X;Montagem\n".encode())
    admin_logado.post(reverse("admin:loja_servico_importar"), {"planilha": arquivo})

    assert ExecucaoIntegracao.objects.get().tipo == ExecucaoIntegracao.Tipo.IMPORTAR_SERVICOS


@pytest.mark.parametrize("referencia", ["MES-1", "mesa-de-centro", "Mesa de Centro",
                                        "  mesa de centro  "])
def test_produto_da_regra_e_achado_pelo_sku_pelo_handle_ou_pelo_titulo(referencia):
    """A planilha vinda da loja traz o handle ou o titulo, nao o SKU do hub."""
    mesa = criar_produto("Mesa de Centro", sku="MES-1", price=Decimal("800.00"))
    mesa.slug = "mesa-de-centro"
    mesa.save()

    servico, _ = importar_linha(_linha(regra_ordem="10", regra_tipo="produto",
                                       produtos=referencia))

    assert list(servico.regras.get().produtos.all()) == [mesa]


def test_titulo_repetido_pede_o_sku_em_vez_de_escolher_um_ao_acaso():
    criar_produto("Poltrona", sku="POL-1", price=Decimal("1.00"))
    criar_produto("Poltrona", sku="POL-2", price=Decimal("1.00"))

    with pytest.raises(LinhaInvalida, match="2 produtos com o titulo 'Poltrona'"):
        importar_linha(_linha(regra_ordem="10", regra_tipo="produto", produtos="Poltrona"))


def test_servico_sem_sku_e_achado_pelo_nome():
    """Planilha com so o nome do servico: acha o cadastro pelo nome e mantem o SKU dele."""
    montagem = Servico.objects.create(nome="Montagem", sku="SERV-MONT")

    servico, criado = importar_linha({"sku": "", "nome": " montagem ", "preco": "90,00"})

    assert not criado and servico.pk == montagem.pk and servico.sku == "SERV-MONT"
    assert servico.preco == Decimal("90.00")


def test_sku_novo_com_nome_que_ja_existe_atualiza_o_sku_do_servico():
    """Criar outro bateria no nome unico; a planilha diz o SKU certo do servico."""
    montagem = Servico.objects.create(nome="Montagem", sku="SERV-MONTAGEM")

    servico, criado = importar_linha({"sku": "SERV-495657", "nome": "Montagem"})

    assert not criado and servico.pk == montagem.pk and servico.sku == "SERV-495657"
