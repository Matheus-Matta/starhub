"""Planilha de frete para importar no Suri (modelo template_politicas.xlsx).

Uma aba por tabela de frete, com o nome dela; uma linha por faixa de CEP.
"""

from decimal import Decimal
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from apps.logistica.models import FaixaCep, TabelaFrete
from apps.suri.frete_planilha import COLUNAS, planilha

pytestmark = pytest.mark.django_db
URL = "admin:integracoes_suri_frete_planilha"


def _tabela(nome, faixas, prazo=3, **extra):
    tabela = TabelaFrete.objects.create(nome=nome, tipo=TabelaFrete.Tipo.FAIXA_CEP,
                                        prazo_dias=prazo, **{"no_checkout": True, **extra})
    for inicial, final, valor, prazo_faixa in faixas:
        FaixaCep.objects.create(tabela=tabela, cep_inicial=inicial, cep_final=final,
                                valor=Decimal(valor), prazo_dias=prazo_faixa)
    return tabela


def _linhas(aba):
    return [[c.value for c in linha] for linha in aba.iter_rows()]


def test_cada_tabela_vira_uma_aba_com_as_faixas_no_formato_do_suri():
    _tabela("Entrega Padrão", [("01000000", "09999999", "50.10", None),
                               ("60000000", "60999999", "12.50", 7)])

    aba = planilha()["Entrega Padrão"]

    assert _linhas(aba) == [
        COLUNAS,
        ["S", "01000000", "09999999", None, None, None, 3, None, None, Decimal("50.10"),
         0, "N", "-1"],
        ["S", "60000000", "60999999", None, None, None, 7, None, None, Decimal("12.50"),
         None, None, None]]
    assert aba["B2"].number_format == "@", "CEP com zero na frente fica texto"


def test_fica_fora_tabela_por_distancia_desativada_ou_fora_do_checkout():
    """O Suri so entende faixa de CEP; tabela que o checkout nao oferece nao vai."""
    _tabela("Vai", [("60000000", "60999999", "10", None)])
    _tabela("Desativada", [("60000000", "60999999", "10", None)], active=False)
    _tabela("Fora do checkout", [("60000000", "60999999", "10", None)], no_checkout=False)
    TabelaFrete.objects.create(nome="Por km", tipo=TabelaFrete.Tipo.DISTANCIA,
                               cep_origem="60000000", preco_por_km=Decimal("2"),
                               no_checkout=True)

    assert planilha().sheetnames == ["Vai"]


def test_nome_de_aba_invalido_no_excel_e_cortado_e_sem_repetir():
    """Excel recusa aba com / ou mais de 31 letras, e duas abas com o mesmo nome."""
    _tabela("Entrega/Expressa " + "x" * 40, [("60000000", "60999999", "10", None)])
    _tabela("Entrega Expressa " + "x" * 40, [("60000000", "60999999", "10", None)])

    nomes = planilha().sheetnames

    assert len(set(nomes)) == 2 and all(len(n) <= 31 and "/" not in n for n in nomes)


def test_sem_tabela_a_planilha_explica_o_que_fazer():
    aba = planilha().active

    assert "oferecer no checkout" in aba["A1"].value


def test_botao_na_secao_de_frete_do_suri_baixa_a_planilha(admin_logado):
    _tabela("Entrega Padrão", [("60000000", "60999999", "10", None)])

    tela = admin_logado.get(reverse("admin:integracoes_configuracaointegracao_suri"))
    resposta = admin_logado.get(reverse(URL))

    assert reverse(URL) in tela.content.decode()
    assert resposta.status_code == 200
    assert "frete-suri" in resposta["Content-Disposition"]
    assert load_workbook(BytesIO(resposta.content)).sheetnames == ["Entrega Padrão"]
