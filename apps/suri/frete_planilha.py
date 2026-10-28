"""Frete do hub em planilha para importar no Portal do Suri (politicas de frete).

Segue o modelo do Suri (template_politicas.xlsx): uma aba por politica, com o nome
dela, e uma linha por faixa de CEP.

    DISPONIVEL | CEP_INICIAL | CEP_FINAL | ... | TEMPO_MAXIMO_DE_ENTREGA | ... | VALOR | ...
    S          | 01000000    | 09999999  |     | 3                       |     | 50.10 |

Vao as tabelas ativas, por faixa de CEP, marcadas "oferecer no checkout": as mesmas
da API de logistica (logistica_api.py). Tabela por distancia fica de fora: a
politica do Suri so compara CEP. FORMATO_TEMPO_ENTREGA, RETIRADA e SELLER_ID sao da
politica inteira e o Suri le so na primeira linha: dias, entrega e "-1" (todas as
lojas). Peso e valor do pedido ficam vazios: o hub nao limita por eles.
"""

import re

from openpyxl import Workbook
from openpyxl.styles import Font

from apps.logistica.models import TabelaFrete

COLUNAS = ["DISPONIVEL", "CEP_INICIAL", "CEP_FINAL", "PESO_MINIMO", "PESO_MAXIMO",
           "TEMPO_MINIMO_DE_ENTREGA", "TEMPO_MAXIMO_DE_ENTREGA", "VALOR_MINIMO_PEDIDO",
           "VALOR_MAXIMO_PEDIDO", "VALOR", "FORMATO_TEMPO_ENTREGA", "RETIRADA", "SELLER_ID"]
EM_DIAS, ENTREGA, TODAS_AS_LOJAS = 0, "N", "-1"
SEM_TABELA = ("Nenhuma tabela por faixa de CEP para exportar. Marque \"oferecer no "
              "checkout\" em Logistica > Tabelas de frete e baixe de novo.")
# Caracteres que o Excel recusa no nome da aba; o nome tem no maximo 31.
PROIBIDOS = re.compile(r"[\\/*?:\[\]]")
LIMITE_NOME = 31


def tabelas():
    return (TabelaFrete.objects.filter(active=True, no_checkout=True,
                                       tipo=TabelaFrete.Tipo.FAIXA_CEP)
            .prefetch_related("faixas").order_by("nome"))


def _nome_da_aba(nome, usados):
    base = PROIBIDOS.sub("-", nome).strip()[:LIMITE_NOME] or "Frete"
    candidato, numero = base, 2
    while candidato.casefold() in usados:  # o Excel compara sem maiuscula
        sufixo = f" ({numero})"
        candidato, numero = base[:LIMITE_NOME - len(sufixo)] + sufixo, numero + 1
    usados.add(candidato.casefold())
    return candidato


def _linhas(tabela):
    for posicao, faixa in enumerate(tabela.faixas.all()):
        prazo = faixa.prazo_dias if faixa.prazo_dias is not None else tabela.prazo_dias
        geral = [EM_DIAS, ENTREGA, TODAS_AS_LOJAS] if posicao == 0 else [None] * 3
        yield ["S", faixa.cep_inicial, faixa.cep_final, None, None, None, prazo, None,
               None, faixa.valor, *geral]


def _aba(livro, titulo, tabela):
    aba = livro.create_sheet(titulo)
    aba.append(COLUNAS)
    for celula in aba[1]:
        celula.font = Font(bold=True)
    for linha in _linhas(tabela):
        aba.append(linha)
    for coluna in ("B", "C"):  # CEP com zero na frente nao pode virar numero
        for celula in aba[coluna]:
            celula.number_format = "@"


def planilha():
    livro = Workbook()
    livro.remove(livro.active)
    usados = set()
    for tabela in tabelas():
        _aba(livro, _nome_da_aba(tabela.nome, usados), tabela)
    if not livro.sheetnames:
        livro.create_sheet("Frete")["A1"] = SEM_TABELA
    return livro
